"""Adapter and document lifecycle checks with mocked retrieval and inference."""
import json
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / 'src'))
import httpx
from docx import Document
from pypdf import PdfWriter
from evidence_api.ingestion import parse_document
from web_demo import semantic
from web_demo.support import parse_upload, retrieve, GroqClient, DemoBudget, answer_document, close_document
from evaluation.run_benchmark import variant_names


class FakeDocument:
    """Test double for plumbing checks; provides no embedding-quality evidence."""
    def __init__(self, name, raw):
        parsed = parse_document(name, raw)
        self.document_id = parsed['document_id']
        self.summary = {'chunks': len(parsed['pages'])}
        self.profile = {'test_double': True}
        self.closed = False
        self.hits = [{'id': f"{self.document_id}_p{p['page']}_c0", 'source': name,
                      'page': p['page'], 'text': p['text'], 'distance': 0.1} for p in parsed['pages']]
    def search(self, question, top_k=5):
        if self.closed: raise ValueError('cleared')
        return self.hits[:top_k]
    def close(self): self.closed = True


class Tests(unittest.TestCase):
    def setUp(self):
        self.mock_session = patch('web_demo.semantic.SemanticDocument', FakeDocument)
        self.mock_session.start(); self.addCleanup(self.mock_session.stop)

    def test_docx_uses_shared_parser(self):
        d = Document(); d.add_paragraph('Domestic approval is optional.')
        table = d.add_table(rows=1, cols=2)
        table.cell(0,0).text = 'International'; table.cell(0,1).text = 'Approval required'
        b = BytesIO(); d.save(b)
        parsed = parse_document('policy.docx', b.getvalue())
        self.assertIn('International | Approval required', parsed['pages'][0]['text'])
        self.assertEqual(parsed['pages'][0]['page'], 0)
        doc = parse_upload('policy.docx', b.getvalue())
        self.assertEqual(doc['document_id'], parsed['document_id'])

    def test_scanned_pdf_rejected(self):
        w = PdfWriter(); w.add_blank_page(width=100, height=100)
        b = BytesIO(); w.write(b)
        with self.assertRaises(ValueError): parse_upload('scan.pdf', b.getvalue())

    def test_document_isolation_and_clear(self):
        a = parse_upload('a.txt', b'Apples have seeds.')
        b = parse_upload('b.txt', b'Bananas are yellow.')
        self.assertTrue(all(h['source']=='a.txt' for h in retrieve(a,'Which apples?')))
        self.assertNotEqual(a['document_id'], b['document_id'])
        close_document(a)
        with self.assertRaises(ValueError): retrieve(a, 'Which apples?')
        self.assertTrue(retrieve(b, 'Which bananas?'))

    def test_invalid_uploads(self):
        for name, raw in [('bad.exe',b'bad'), ('empty.txt',b''), ('big.txt',b'x'*(5*1024*1024+1))]:
            with self.assertRaises(ValueError): parse_upload(name,raw)

    def test_budget(self):
        b=DemoBudget(limit=2,window=100)
        self.assertTrue(b.reserve(0)); self.assertTrue(b.reserve(1))
        self.assertFalse(b.reserve(2)); self.assertTrue(b.reserve(101))

    def test_quota_stops_without_retry(self):
        calls=[]
        def handler(request): calls.append(request); return httpx.Response(429)
        c=GroqClient('TEST_SECRET','openai/gpt-oss-20b',httpx.MockTransport(handler))
        with self.assertRaises(RuntimeError):
            c.chat(model=c.model,messages=[{'role':'system','content':'JSON'}],
                   format={'type':'object','properties':{},'required':[],'additionalProperties':False},
                   options={},think=False,stream=False)
        self.assertEqual(len(calls),1)

    def test_provider_bundle_and_metadata(self):
        doc=parse_upload('policy.txt',b'Domestic travel does not require advance approval.')
        outputs=[{'answerability':'supported','evidence_handles':['S1']},
                 {'response_type':'supported','claims':[{'text':'Domestic travel does not require advance approval.','evidence':['E1']}]},
                 {'answers_question':True,'coverage_reason':'Answers it.','abstention':'not_applicable',
                  'claims':[{'claim_index':0,'verdict':'supported','reason':'Explicit.'}]}]
        calls=[]
        def handler(request):
            calls.append(json.loads(request.content))
            return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(outputs[len(calls)-1])}}],
                                           'usage':{'completion_tokens':10}})
        result,bundle,_=answer_document(doc,'Does domestic travel require approval?','openai/gpt-oss-20b',
                                       'TEST_SECRET',transport=httpx.MockTransport(handler))
        self.assertEqual(result['status'],'needs_factual_review')
        self.assertEqual(result['demo_profile']['retrieval'],semantic.PROFILE)
        self.assertEqual(result['demo_profile']['top_k'],5)
        self.assertEqual(calls[0]['max_completion_tokens'],4096)
        import zipfile
        with zipfile.ZipFile(BytesIO(bundle)) as archive:
            for name in archive.namelist():
                self.assertNotIn(b'TEST_SECRET',archive.read(name))
                self.assertNotIn(b'_semantic_session',archive.read(name))

    def test_matrix_unchanged(self):
        plan=json.loads((ROOT/'evaluation/cases.json').read_text())
        self.assertEqual(len(plan['cases']),8)
        self.assertEqual(sum(len(variant_names(c)) for c in plan['cases'])*3,66)


class SharedFactoryTests(unittest.TestCase):
    def test_assets_call_local_engine_with_isolated_root(self):
        semantic.retrieval_assets.cache_clear()
        calls=[]
        def build_engine(**kwargs):
            calls.append(kwargs)
            return SimpleNamespace(store=SimpleNamespace(encoder='encoder',splitter='splitter',profile={'same':True}))
        with patch.dict(sys.modules,{'evidence_api.app':SimpleNamespace(build_engine=build_engine)}):
            encoder,splitter,profile,directory=semantic.retrieval_assets()
            try:
                self.assertEqual((encoder,splitter),('encoder','splitter'))
                self.assertEqual(profile,{'same':True})
                self.assertEqual(calls,[{'data_root':directory.name}])
            finally:
                directory.cleanup(); semantic.retrieval_assets.cache_clear()

class SemanticLifecycleTests(unittest.TestCase):
    def test_store_filter_and_cleanup(self):
        paths=[]; searches=[]; deleted=[]
        def factory(path):
            paths.append(path)
            def search(question, top_k, ids):
                searches.append((question, top_k, ids)); return []
            return SimpleNamespace(
                ingest=lambda name,raw:{'document_id':'doc_test','chunks':1},
                profile={'test':True},search=search,
                client=SimpleNamespace(delete_collection=lambda name:deleted.append(name)),
                collection=SimpleNamespace(name='session_collection'))
        a=semantic.SemanticDocument('a.txt',b'text',store_factory=factory)
        b=semantic.SemanticDocument('b.txt',b'text',store_factory=factory)
        self.assertNotEqual(paths[0],paths[1])
        a.search('question')
        self.assertEqual(searches,[('question',5,['doc_test'])])
        a.close(); a.close(); b.close()
        self.assertTrue(all(not Path(path).exists() for path in paths))
        self.assertEqual(deleted,['session_collection','session_collection'])
        with self.assertRaises(ValueError): a.search('question')

    def test_failed_factory_cleans_temporary_directory(self):
        paths=[]
        def factory(path):
            paths.append(path); raise RuntimeError('load failed')
        with self.assertRaises(RuntimeError):
            semantic.SemanticDocument('a.txt',b'text',store_factory=factory)
        self.assertFalse(Path(paths[0]).exists())


if __name__=='__main__': unittest.main()
